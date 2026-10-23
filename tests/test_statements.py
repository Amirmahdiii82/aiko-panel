"""Financial statements: the balance sheet balances, profit matches the reports, cash ties opening to closing."""
import io

from app.modules.erp import db, ledger

from .conftest import H

OWNER = {"id": None, "name": "امیر"}


def _accounts():
    with db.read() as con:
        acc = {r["name"]: r["id"] for r in con.execute("SELECT id, name FROM accounts")}
        partners = [r[0] for r in con.execute("SELECT id FROM partners ORDER BY id")]
        cards = {r[0]: r[1] for r in con.execute("SELECT partner_id, id FROM accounts WHERE kind = 'partner'")}
    return acc, partners, cards


def _month(owner):
    acc, partners, cards = _accounts()
    bank, cash = acc["حساب بانکی فروشگاه"], acc["صندوق نقدی"]
    with db.tx() as con:
        ledger.save_capital(con, {"date": "2026-10-01", "partner_id": partners[0], "kind": "in", "amount": 10_000_000,
                                  "account_id": bank}, OWNER)
        pid = ledger.create_product(con, {"name": "رژ لب", "price": 300_000}, None, quiet=True)
        ledger.create_purchase(con, {"date": "2026-10-02", "account_id": bank,
                                     "items": [{"product_id": pid, "qty": 10, "unit_cost": 200_000}]}, None, quiet=True)
        ledger.create_sale(con, {"date": "2026-10-03", "channel": "in_person", "account_id": cash,
                                 "items": [{"product_id": pid, "qty": 3, "unit_price": 300_000}]}, None, quiet=True)
        ledger.create_sale(con, {"date": "2026-10-04", "channel": "instagram", "paid": False,
                                 "items": [{"product_id": pid, "qty": 1, "unit_price": 300_000}]}, None, quiet=True)
        ledger.save_expense(con, {"date": "2026-10-05", "category": "تبلیغات", "amount": 100_000,
                                  "account_id": cards[partners[1]]}, OWNER)       # paid from a partner's own card
    return owner.get(f"/api/erp/statements?period=custom&start=2026-10-01&end={db.today()}").json()


def amount(rows, label):
    return next(r["amount"] for r in rows if r["label"] == label)


def test_the_balance_sheet_balances(owner):
    s = _month(owner)
    assets, right = s["balance"]["assets"], s["balance"]["liabilities_equity"]
    assert amount(assets, "حساب بانکی فروشگاه") == 8_000_000
    assert amount(assets, "صندوق نقدی") == 900_000
    assert amount(assets, "طلب از مشتری‌ها (نسیه)") == 300_000
    assert amount(assets, "موجودی کالا") == 6 * 200_000
    assert amount(assets, "جمع دارایی‌ها") == amount(right, "جمع بدهی‌ها و حقوق صاحبان سرمایه") == 10_400_000
    assert amount(right, "سود (زیان) انباشته") == 300_000
    assert not any(r["label"] == "موجودی اولیه و اصلاحات انبار" for r in right)    # everything is explained


def test_the_income_statement(owner):
    rows = _month(owner)["income"]["rows"]
    assert amount(rows, "درآمد خالص") == 1_200_000
    assert amount(rows, "بهای تمام‌شده کالای فروش‌رفته") == -800_000
    assert amount(rows, "سود ناخالص") == 400_000
    assert amount(rows, "تبلیغات") == -100_000
    assert amount(rows, "سود (زیان) خالص") == 300_000


def test_cash_flow_ties_opening_to_closing(owner):
    rows = _month(owner)["cash"]["rows"]
    assert amount(rows, "نقد ابتدای دوره") == 0
    assert amount(rows, "دریافت از فروش") == 900_000
    assert amount(rows, "پرداخت برای خرید کالا") == -2_000_000
    assert amount(rows, "پرداخت هزینه‌ها") == 0                 # paid from a partner's card, not the shop's cash
    assert amount(rows, "آورده شرکا") == 10_000_000
    assert amount(rows, "نقد پایان دوره") == amount(rows, "نقد ابتدای دوره") + amount(rows, "خالص تغییر نقد") == 8_900_000


def test_opening_stock_without_a_payment_is_one_explained_line(owner):
    with db.tx() as con:
        ledger.create_product(con, {"name": "کالای موجودی اولیه", "stock": 5, "avg_cost": 100_000}, None, quiet=True)
    s = owner.get(f"/api/erp/statements?period=all").json()
    right = s["balance"]["liabilities_equity"]
    assert amount(right, "موجودی اولیه و اصلاحات انبار") == 500_000
    assert amount(s["balance"]["assets"], "جمع دارایی‌ها") == amount(right, "جمع بدهی‌ها و حقوق صاحبان سرمایه")


def test_the_three_statements_in_excel(owner):
    import openpyxl
    _month(owner)
    r = owner.get(f"/api/erp/statements.xlsx?period=custom&start=2026-10-01&end={db.today()}")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["ترازنامه", "سود و زیان", "جریان نقدی"]
    values = {row[0].strip(): row[1] for row in wb["سود و زیان"].iter_rows(min_row=2, values_only=True) if row and row[0]}
    assert values["سود (زیان) خالص"] == 300_000
