"""Financial statements, from the same ledger as everything else (no separate bookkeeping to keep in step):

* Balance sheet (ترازنامه) at a date: what the shop has (cash in its accounts, money with sellers, unpaid sales, stock
  at average cost) against what it owes and what belongs to the partners.
* Income statement (سود و زیان) for a period: net sales → cost of goods sold → gross profit → running expenses by
  category → net profit. Same rules as the reports page (stock bought is not an expense; it becomes cost when sold).
* Cash flow (جریان نقدی) for a period: money in and out of the shop's own accounts and sellers' hands, from opening
  to closing balance.

Partners' personal cards are not the shop's cash: what a partner paid from their card counts as their contribution,
shop money they hold counts against it (the same "net stake" as the partners page). The panel is not double-entry, so
the balance sheet closes with one explained line: goods that came into stock without a recorded payment (opening
stock, products created from the site) minus stock written off.

Every statement is a list of rows {label, amount, kind: section | line | subtotal | total, indent, hint}.
"""
import io

from ...core import jalali
from . import db, ledger

CASH_KINDS = ("shop", "agent")


def _row(label: str, amount: int | None = None, kind: str = "line", indent: int = 0, hint: str = "") -> dict:
    return {"label": label, "amount": None if amount is None else int(round(amount)), "kind": kind, "indent": indent,
            "hint": hint}


def _profit_to(con, until: str) -> float:
    r = con.execute("""SELECT IFNULL(SUM(total), 0), IFNULL(SUM(cogs), 0) FROM sales
                       WHERE status = 'done' AND date <= ?""", (until,)).fetchone()
    expenses = con.execute("SELECT IFNULL(SUM(amount), 0) FROM expenses WHERE is_inventory = 0 AND date <= ?",
                           (until,)).fetchone()[0]
    return r[0] - r[1] - expenses


def _stock_value(con, until: str) -> tuple[int, int]:
    """(units, value at the current average cost) on that date."""
    if until >= db.today():
        r = con.execute("""SELECT IFNULL(SUM(MAX(stock, 0)), 0), IFNULL(SUM(MAX(stock, 0) * avg_cost), 0)
                           FROM products""").fetchone()
        return int(r[0]), int(round(r[1]))
    r = con.execute("""SELECT IFNULL(SUM(MAX(q, 0)), 0), IFNULL(SUM(MAX(q, 0) * avg_cost), 0) FROM (
                         SELECT p.avg_cost, IFNULL(SUM(m.qty), 0) AS q FROM products p
                         LEFT JOIN stock_moves m ON m.product_id = p.id AND substr(m.created_at, 1, 10) <= ?
                         GROUP BY p.id)""", (until,)).fetchone()
    return int(r[0]), int(round(r[1]))


def balance_sheet(con, as_of: str) -> dict:
    balances = ledger.account_balances(con, as_of)
    shop = [b for b in balances if b["kind"] == "shop" and (b["balance"] or b["active"])]
    agents = [b for b in balances if b["kind"] == "agent" and b["balance"]]
    receivable = con.execute("""SELECT IFNULL(SUM(total), 0), COUNT(*) FROM sales
                                WHERE status = 'done' AND paid = 0 AND date <= ?""", (as_of,)).fetchone()
    units, stock = _stock_value(con, as_of)

    assets = [_row("دارایی‌های جاری", kind="section")]
    assets += [_row(b["name"], b["balance"], indent=1) for b in shop]
    held = sum(b["balance"] for b in agents if b["balance"] > 0)
    if held:
        assets.append(_row("پول نزد فروشنده‌ها", held, indent=1, hint="فروش‌هایی که فروشنده‌ها انجام داده‌اند و پولش را هنوز تحویل نداده‌اند"))
    if receivable[0]:
        assets.append(_row("طلب از مشتری‌ها (نسیه)", receivable[0], indent=1, hint=f"{ledger.fa(receivable[1])} فروش پرداخت‌نشده"))
    assets.append(_row("موجودی کالا", stock, indent=1, hint=f"{ledger.fa(units)} عدد، به میانگین قیمت تمام‌شده"))
    total_assets = sum(b["balance"] for b in shop) + held + receivable[0] + stock
    assets.append(_row("جمع دارایی‌ها", total_assets, "total"))

    owed_sellers = -sum(b["balance"] for b in agents if b["balance"] < 0)
    liabilities = [_row("بدهی‌ها", kind="section")]
    if owed_sellers:
        liabilities.append(_row("بدهی به فروشنده‌ها (پورسانت)", owed_sellers, indent=1))
    else:
        liabilities.append(_row("بدهی ثبت‌شده‌ای نیست", 0, indent=1))
    liabilities.append(_row("جمع بدهی‌ها", owed_sellers, "subtotal"))

    partners = db.rows(con.execute("SELECT id, name FROM partners WHERE active = 1 ORDER BY id"))
    cap = {r[0]: r[1] for r in con.execute(
        """SELECT partner_id, SUM(CASE WHEN kind = 'in' THEN amount ELSE -amount END) FROM capital
           WHERE date <= ? GROUP BY partner_id""", (as_of,))}
    stakes = []
    for p in partners:
        held_by_partner = sum(b["balance"] for b in balances if b["kind"] == "partner" and b["partner_id"] == p["id"])
        stakes.append((p["name"], (cap.get(p["id"]) or 0) - held_by_partner))
    profit = _profit_to(con, as_of)
    equity_known = sum(s for _, s in stakes) + profit
    adjustment = total_assets - owed_sellers - equity_known
    equity = [_row("حقوق صاحبان سرمایه", kind="section")]
    equity += [_row(f"آورده خالص {name}", stake, indent=1, hint="آورده و پرداخت از کارت شخصی، منهای برداشت" if i == 0 else "")
               for i, (name, stake) in enumerate(stakes)]
    equity.append(_row("سود (زیان) انباشته", profit, indent=1, hint="از ابتدا تا این تاریخ"))
    if abs(adjustment) >= 1:
        equity.append(_row("موجودی اولیه و اصلاحات انبار", adjustment, indent=1,
                           hint="کالای واردشده بدون پرداخت ثبت‌شده، منهای کسری انبار"))
    equity.append(_row("جمع حقوق صاحبان سرمایه", equity_known + adjustment, "subtotal"))
    right = liabilities + equity + [_row("جمع بدهی‌ها و حقوق صاحبان سرمایه", owed_sellers + equity_known + adjustment, "total")]
    return {"as_of": as_of, "assets": assets, "liabilities_equity": right}


def income_statement(con, start: str, end: str) -> dict:
    s = con.execute("""SELECT COUNT(*), IFNULL(SUM(subtotal), 0), IFNULL(SUM(discount), 0), IFNULL(SUM(shipping), 0),
                              IFNULL(SUM(fees), 0), IFNULL(SUM(tax), 0), IFNULL(SUM(refunded), 0), IFNULL(SUM(total), 0),
                              IFNULL(SUM(cogs), 0)
                       FROM sales WHERE status = 'done' AND date BETWEEN ? AND ?""", (start, end)).fetchone()
    orders, gross_sales, discount, shipping, fees, tax, refunded, net_sales, cogs = s
    other = net_sales - (gross_sales - discount + shipping + fees + tax - refunded)
    expenses = db.rows(con.execute("""SELECT category, SUM(amount) AS amount FROM expenses
                                      WHERE is_inventory = 0 AND date BETWEEN ? AND ?
                                      GROUP BY category ORDER BY amount DESC""", (start, end)))
    rows = [_row("درآمد", kind="section"),
            _row("فروش کالا", gross_sales, indent=1, hint=f"{ledger.fa(orders)} فروش")]
    if discount:
        rows.append(_row("تخفیف‌ها", -discount, indent=1))
    if refunded:
        rows.append(_row("برگشت از فروش (سایت)", -refunded, indent=1))
    if shipping:
        rows.append(_row("هزینه ارسال دریافتی از مشتری", shipping, indent=1))
    if fees or tax or abs(other) >= 1:
        rows.append(_row("کارمزد، مالیات و سایر", fees + tax + other, indent=1))
    rows += [_row("درآمد خالص", net_sales, "subtotal"),
             _row("بهای تمام‌شده کالای فروش‌رفته", -cogs, indent=1, hint="به میانگین قیمت خرید (با هزینه حمل) هنگام فروش"),
             _row("سود ناخالص", net_sales - cogs, "subtotal",
                  hint=f"حاشیه {ledger.fa(round((net_sales - cogs) / net_sales * 100, 1)) if net_sales else '۰'}٪"),
             _row("هزینه‌های جاری", kind="section")]
    rows += [_row(e["category"], -e["amount"], indent=1) for e in expenses] or [_row("هزینه‌ای ثبت نشده", 0, indent=1)]
    total_exp = sum(e["amount"] for e in expenses)
    rows += [_row("جمع هزینه‌ها", -total_exp, "subtotal"),
             _row("سود (زیان) خالص", net_sales - cogs - total_exp, "total")]
    missing = ledger.missing_costs(con, start, end)
    return {"rows": rows, "missing_cost": missing}


def cash_flow(con, start: str, end: str) -> dict:
    """Money of the shop's own accounts and in sellers' hands (partners' personal cards are not the shop's cash)."""
    cash_ids = [r[0] for r in con.execute(f"SELECT id FROM accounts WHERE kind IN {CASH_KINDS}")]
    marks = ",".join("?" * len(cash_ids)) or "NULL"

    def total(where: str, args: tuple) -> int:
        return int(con.execute(f"SELECT IFNULL(SUM(amount), 0) FROM money_flows WHERE account_id IN ({marks}) AND {where}",
                               (*cash_ids, *args)).fetchone()[0])
    opening = total("date < ?", (start,))
    closing = total("date <= ?", (end,))
    by_src = {r[0]: (int(r[1] or 0), int(r[2] or 0)) for r in con.execute(
        f"""SELECT src, SUM(CASE WHEN amount > 0 THEN amount ELSE 0 END), SUM(CASE WHEN amount < 0 THEN amount ELSE 0 END)
            FROM money_flows WHERE account_id IN ({marks}) AND date BETWEEN ? AND ? AND src != 'transfer' GROUP BY src""",
        (*cash_ids, start, end))}
    t_in = t_out = 0
    for t in con.execute("SELECT from_account, to_account, amount FROM transfers WHERE date BETWEEN ? AND ?", (start, end)):
        f_cash, t_cash = t["from_account"] in cash_ids, t["to_account"] in cash_ids
        if t_cash and not f_cash:
            t_in += t["amount"]
        elif f_cash and not t_cash:
            t_out -= t["amount"]
    sale_in = by_src.get("sale", (0, 0))[0]
    purchase_out = by_src.get("purchase", (0, 0))[1]
    expense_out = sum(by_src.get("expense", (0, 0)))
    cap_in, cap_out = by_src.get("capital", (0, 0))
    operating = sale_in + purchase_out + expense_out
    financing = cap_in + cap_out + t_in + t_out
    rows = [_row("نقد ابتدای دوره", opening, "opening"),
            _row("فعالیت‌های عملیاتی", kind="section"),
            _row("دریافت از فروش", sale_in, indent=1),
            _row("پرداخت برای خرید کالا", purchase_out, indent=1),
            _row("پرداخت هزینه‌ها", expense_out, indent=1, hint="شامل پورسانت‌ها"),
            _row("خالص جریان عملیاتی", operating, "subtotal"),
            _row("شرکا", kind="section"),
            _row("آورده شرکا", cap_in, indent=1),
            _row("برداشت شرکا", cap_out, indent=1)]
    if t_in or t_out:
        rows.append(_row("انتقال با کارت شخصی شرکا", t_in + t_out, indent=1))
    rows += [_row("خالص جریان شرکا", financing, "subtotal"),
             _row("خالص تغییر نقد", operating + financing, "subtotal"),
             _row("نقد پایان دوره", closing, "total")]
    return {"rows": rows}


def statements(con, start: str, end: str, label: str) -> dict:
    as_of = min(end, db.today())
    jy, jm, jd = jalali.iso_to_jalali(as_of)
    return {"label": label, "start": start, "end": end, "as_of": as_of,
            "as_of_label": f"{ledger.fa(jd)} {jalali.MONTHS[jm - 1]} {ledger.fa(jy)}",
            "balance": balance_sheet(con, as_of), "income": income_statement(con, start, end),
            "cash": cash_flow(con, max(start, "2000-01-01"), as_of)}


def workbook_bytes(data: dict) -> bytes:
    """The three statements in one Excel file, laid out like the page."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    thin, double = Side(style="thin", color="999999"), Side(style="double", color="333333")

    def sheet(title: str, heading: str, blocks: list[list[dict]]):
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 48
        ws.append([heading])
        ws["A1"].font = Font(bold=True, size=13)
        for rows in blocks:
            ws.append([])
            for r in rows:
                ws.append(["   " * r["indent"] + r["label"], r["amount"], r["hint"] or None])
                row = ws.max_row
                ws.cell(row, 2).number_format = '#,##0;(#,##0)'
                ws.cell(row, 2).alignment = Alignment(horizontal="left")
                ws.cell(row, 3).font = Font(color="888888", size=9)
                if r["kind"] in ("section", "subtotal", "total"):
                    ws.cell(row, 1).font = Font(bold=True)
                    ws.cell(row, 2).font = Font(bold=True)
                if r["kind"] == "subtotal":
                    ws.cell(row, 2).border = Border(top=thin)
                if r["kind"] == "total":
                    ws.cell(row, 2).border = Border(top=thin, bottom=double)

    b = data["balance"]
    sheet("ترازنامه", f"ترازنامه در {data['as_of_label']} (تومان)", [b["assets"], b["liabilities_equity"]])
    sheet("سود و زیان", f"صورت سود و زیان — {data['label']} (تومان)", [data["income"]["rows"]])
    sheet("جریان نقدی", f"صورت جریان نقدی — {data['label']} (تومان)", [data["cash"]["rows"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
