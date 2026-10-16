"""Excel export of everything (one sheet per table, right-to-left, Persian headers, Jalali dates)."""
import io

from . import db, jalali, ledger


def _jdate(iso: str | None) -> str:
    if not iso:
        return ""
    jy, jm, jd = jalali.iso_to_jalali(iso)
    return f"{jy}/{jm:02d}/{jd:02d}"


SHEETS = [
    ("محصولات", """SELECT sku, name, brand, category, price, ROUND(avg_cost), stock, ROUND(MAX(stock, 0) * avg_cost),
                         low_stock, CASE WHEN woo_id IS NULL THEN '' ELSE woo_id || IFNULL('/' || woo_variation_id, '') END,
                         CASE active WHEN 1 THEN 'فعال' ELSE 'بایگانی' END
                  FROM products ORDER BY active DESC, name""",
     ["کد", "نام", "برند", "دسته", "قیمت فروش", "میانگین قیمت تمام‌شده", "موجودی", "ارزش موجودی", "حد هشدار",
      "شناسه سایت", "وضعیت"], []),
    ("فروش", """SELECT s.code, s.date, s.channel, IFNULL(u.name, ''), IFNULL(s.customer_name, ''), IFNULL(s.customer_phone, ''),
                      s.subtotal, s.discount, s.shipping, s.total, ROUND(s.cogs), ROUND(s.total - s.cogs), s.commission,
                      CASE s.paid WHEN 1 THEN 'بله' ELSE 'خیر' END, IFNULL(a.name, ''),
                      CASE s.status WHEN 'done' THEN '' ELSE 'باطل' END, IFNULL(s.note, '')
               FROM sales s LEFT JOIN users u ON u.id = s.seller_id LEFT JOIN accounts a ON a.id = s.account_id
               ORDER BY s.date, s.id""",
     ["کد", "تاریخ", "کانال", "فروشنده", "مشتری", "تلفن", "جمع اقلام", "تخفیف", "ارسال", "مبلغ کل", "بهای تمام‌شده",
      "سود ناخالص", "پورسانت", "پرداخت شده", "حساب دریافت", "وضعیت", "یادداشت"], [1]),
    ("اقلام فروش", """SELECT s.code, s.date, IFNULL(p.sku, ''), i.name, i.qty, i.unit_price, ROUND(i.unit_cost), i.line_total,
                            ROUND(i.line_total - i.qty * i.unit_cost), CASE s.status WHEN 'done' THEN '' ELSE 'باطل' END
                     FROM sale_items i JOIN sales s ON s.id = i.sale_id LEFT JOIN products p ON p.id = i.product_id
                     ORDER BY s.date, s.id""",
     ["کد فروش", "تاریخ", "کد کالا", "کالا", "تعداد", "قیمت واحد", "قیمت تمام‌شده واحد", "جمع", "سود", "وضعیت"], [1]),
    ("خرید", """SELECT p.code, p.date, IFNULL(p.supplier, ''), pr.sku, pr.name, i.qty, ROUND(i.unit_cost), ROUND(i.landed_cost),
                      i.line_total, IFNULL(a.name, ''), CASE p.status WHEN 'done' THEN '' ELSE 'باطل' END
               FROM purchase_items i JOIN purchases p ON p.id = i.purchase_id JOIN products pr ON pr.id = i.product_id
               LEFT JOIN accounts a ON a.id = p.account_id ORDER BY p.date, p.id""",
     ["کد", "تاریخ", "تأمین‌کننده", "کد کالا", "کالا", "تعداد", "قیمت خرید", "قیمت تمام‌شده (با حمل)", "جمع",
      "پرداخت از", "وضعیت"], [1]),
    ("هزینه‌ها", """SELECT e.code, e.date, e.category, IFNULL(e.description, ''), e.amount, IFNULL(a.name, ''),
                          CASE e.is_inventory WHEN 1 THEN 'خرید کالا' ELSE 'هزینه جاری' END, IFNULL(e.note, '')
                   FROM expenses e LEFT JOIN accounts a ON a.id = e.account_id ORDER BY e.date, e.id""",
     ["کد", "تاریخ", "دسته", "شرح", "مبلغ", "پرداخت از", "نوع", "یادداشت"], [1]),
    ("آورده و برداشت", """SELECT c.date, p.name, CASE c.kind WHEN 'in' THEN 'آورده' ELSE 'برداشت' END, c.amount,
                                IFNULL(a.name, ''), IFNULL(c.note, '')
                         FROM capital c JOIN partners p ON p.id = c.partner_id LEFT JOIN accounts a ON a.id = c.account_id
                         ORDER BY c.date""",
     ["تاریخ", "شریک", "نوع", "مبلغ", "حساب", "یادداشت"], [0]),
    ("انتقال‌ها", """SELECT t.date, fa.name, ta.name, t.amount, t.kind, IFNULL(t.note, '')
                    FROM transfers t JOIN accounts fa ON fa.id = t.from_account JOIN accounts ta ON ta.id = t.to_account
                    ORDER BY t.date""",
     ["تاریخ", "از", "به", "مبلغ", "نوع", "یادداشت"], [0]),
]


def workbook_bytes() -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="EEECF8")
    with db.read() as con:
        for title, sql, headers, date_cols in SHEETS:
            ws = wb.create_sheet(title)
            ws.sheet_view.rightToLeft = True
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            for row in con.execute(sql):
                row = list(row)
                for c in date_cols:
                    row[c] = _jdate(row[c])
                if title == "فروش":
                    row[2] = ledger.CHANNELS.get(row[2], row[2])
                ws.append(row)
            for i, h in enumerate(headers, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(40, len(h) + 8))
            ws.freeze_panes = "A2"
        ws = wb.create_sheet("حساب‌ها و شرکا")
        ws.sheet_view.rightToLeft = True
        ws.append(["حساب", "نوع", "مانده"])
        for a in ledger.account_balances(con):
            ws.append([a["name"], {"shop": "فروشگاه", "partner": "شریک", "agent": "فروشنده"}[a["kind"]], a["balance"]])
        ws.append([])
        ws.append(["شریک", "سهم ٪", "آورده خالص", "سهم منصفانه", "اختلاف (+ طلبکار / − بدهکار)"])
        for p in ledger.partner_positions(con)["partners"]:
            ws.append([p["name"], p["share"], p["net"], round(p["fair"]), p["excess"]])
        ws.column_dimensions["A"].width = 28
        for col in "BCDE":
            ws.column_dimensions[col].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
