"""Stock count with an Excel sheet: download every product (warehouse and site, with both current numbers and an
empty "counted" column), fill in what is on the shelf, upload it. The panel takes the counted numbers and the site
follows the panel (stocksync.py), so one count sets both. Empty rows are left as they are.
"""
import io
import re

from . import db, ledger, notify, stocksync

TITLE = "شمارش موجودی"
HEADERS = ["شناسه (تغییر ندهید)", "کد", "نام محصول", "برند", "دسته", "در سایت", "موجودی فعلی پنل", "موجودی فعلی سایت",
           "موجودی شمرده‌شده (پر کنید)"]
COUNT_COL = len(HEADERS)            # the column people fill in (1-based)
DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def _site_text(p: dict) -> str:
    if not p["woo_id"]:
        return ""
    if p["woo_state"] == "nocount":
        return "موجود، بدون عدد" if p["woo_in_stock"] else "ناموجود، بدون عدد"
    if p["woo_state"] == "whole":
        return f"{p['woo_stock']} (همه رنگ‌ها)" if p["woo_stock"] is not None else "همه رنگ‌ها"
    return p["woo_stock"] if p["woo_stock"] is not None else "؟"


def workbook_bytes() -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    with db.read() as con:
        rows = db.rows(con.execute(
            """SELECT id, sku, name, brand, category, stock, woo_id, woo_state, woo_stock, woo_in_stock FROM products
               WHERE active = 1 ORDER BY CASE WHEN woo_id IS NULL THEN 1 ELSE 0 END, IFNULL(category, 'ي'), name"""))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TITLE
    ws.sheet_view.rightToLeft = True
    ws.append(HEADERS)
    head_fill, fill_me, quiet = PatternFill("solid", fgColor="EEECF8"), PatternFill("solid", fgColor="FFF4CC"), Font(color="999999")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(1, COUNT_COL).fill = PatternFill("solid", fgColor="F4C542")
    for p in rows:
        ws.append([p["id"], p["sku"] or "", p["name"], p["brand"] or "", p["category"] or "",
                   "بله" if p["woo_id"] else "فقط انبار", p["stock"], _site_text(p), None])
        r = ws.max_row
        ws.cell(r, 1).font = quiet
        ws.cell(r, COUNT_COL).fill = fill_me
    rule = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                          error="فقط عدد صحیح صفر یا بیشتر", errorTitle="تعداد نامعتبر")
    ws.add_data_validation(rule)
    if rows:
        rule.add(f"{openpyxl.utils.get_column_letter(COUNT_COL)}2:{openpyxl.utils.get_column_letter(COUNT_COL)}{len(rows) + 1}")
    for i, width in enumerate([10, 12, 52, 16, 14, 11, 14, 16, 20], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "C2"

    guide = wb.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 100
    for line in [
        "شمارش موجودی — پنل و سایت با هم",
        "",
        "۱. تعداد واقعی هر محصول را در ستون زرد «موجودی شمرده‌شده» بنویسید (اگر نیست، ۰).",
        "۲. ستون «شناسه» را تغییر ندهید و ردیف اضافه نکنید؛ ترتیب ردیف‌ها مهم نیست.",
        "۳. ردیفی که خالی بماند دست نمی‌خورد: عدد فعلی پنل می‌ماند و سایت هم همان را نشان می‌دهد.",
        "۴. فایل را در «فروشگاه ← اتصال سایت ← بارگذاری فایل شمارش» بارگذاری کنید.",
        "",
        "بعد از بارگذاری: موجودی پنل همان عدد شمرده‌شده می‌شود و چند ثانیه بعد سایت هم همان عدد را نشان می‌دهد،",
        "حتی محصولاتی که در سایت «مدیریت موجودی» خاموش بود.",
        "",
        "اگر بین دانلود و بارگذاری فروشی ثبت شد، بعد از بارگذاری آن محصول را در انبار اصلاح کنید.",
    ]:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _number(v) -> int | None:
    """A counted number from a cell: 12, 12.0, "۱۲", " 12 ". Empty → None. Anything else is an error."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if v != int(v) or v < 0:
            raise ValueError
        return int(v)
    text = str(v).translate(DIGITS).strip().replace(",", "").replace("٬", "")
    if not re.fullmatch(r"\d+(\.0+)?", text):
        raise ValueError
    return int(float(text))


def read_counts(data: bytes) -> tuple[dict[int, int], list[str]]:
    """{product id: counted number} from an uploaded sheet, and the problems found (shown to the person)."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise ledger.LedgerError("فایل اکسل خوانده نشد؛ همان فایلی که از پنل گرفتید را با فرمت xlsx ذخیره کنید")
    ws = wb[TITLE] if TITLE in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(rows, [])]
    try:
        id_col = next(i for i, h in enumerate(header) if h.startswith("شناسه"))
        count_col = next(i for i, h in enumerate(header) if h.startswith("موجودی شمرده"))
    except StopIteration:
        raise ledger.LedgerError("ستون «شناسه» یا «موجودی شمرده‌شده» پیدا نشد؛ از فایلی که پنل می‌دهد استفاده کنید")
    sku_col = header.index("کد") if "کد" in header else None
    name_col = header.index("نام محصول") if "نام محصول" in header else None
    with db.read() as con:
        by_sku = {r[0]: r[1] for r in con.execute("SELECT sku, id FROM products WHERE active = 1 AND sku IS NOT NULL")}
        ids = {r[0] for r in con.execute("SELECT id FROM products WHERE active = 1")}
    counts, problems = {}, []
    for n, row in enumerate(rows, 2):
        if not row or all(c in (None, "") for c in row):
            continue
        name = row[name_col] if name_col is not None and name_col < len(row) else f"ردیف {n}"
        try:
            qty = _number(row[count_col] if count_col < len(row) else None)
        except ValueError:
            problems.append(f"ردیف {ledger.fa(n)} ({name}): «{row[count_col]}» عدد درستی نیست")
            continue
        if qty is None:
            continue
        pid = row[id_col] if id_col < len(row) else None
        try:
            pid = int(str(pid).translate(DIGITS)) if pid not in (None, "") else None
        except ValueError:
            pid = None
        if pid not in ids and sku_col is not None and sku_col < len(row):
            pid = by_sku.get(str(row[sku_col] or "").strip())
        if pid not in ids:
            problems.append(f"ردیف {ledger.fa(n)} ({name}): این محصول در پنل پیدا نشد")
            continue
        counts[pid] = qty
    if not counts and not problems:
        raise ledger.LedgerError("هیچ عددی در ستون «موجودی شمرده‌شده» نوشته نشده")
    return counts, problems


def apply(counts: dict[int, int], user: dict) -> dict:
    """The panel takes the counted numbers; every counted product that is on the site is sent to it (and starts being
    counted there if it was not). Stock sync starts if it had not yet."""
    changed = same = 0
    with db.tx() as con:
        for pid, qty in counts.items():
            p = ledger.get_product(con, pid)
            if qty != p["stock"]:
                changed += 1
                ledger.add_move(con, pid, qty - p["stock"], "adjust", unit_cost=p["avg_cost"] if qty > p["stock"] else None,
                                note="شمارش انبار (فایل اکسل)", user_id=user["id"], quiet=True)
            else:
                same += 1
            if p["woo_id"] and p["woo_state"] not in ("whole", "gone"):
                con.execute("UPDATE products SET woo_counted = 1, woo_state = 'pending' WHERE id = ?", (pid,))
        s = db.get_setting("woo", {}, con) or {}
        if not s.get("stock_start"):
            s.update(stock_start="count", stock_started_at=db.now(), stock_started_by=user["name"])
            db.set_setting("woo", s, con)
        notify.activity(con, "stock_count", f"📋 شمارش انبار ثبت شد: {ledger.fa(len(counts))} محصول",
                        f"{ledger.fa(changed)} محصول عوض شد، {ledger.fa(same)} همان بود. سایت هم همین اعداد را می‌گیرد. — {user['name']}",
                        user_id=user["id"], level="success", link="#/erp/site")
        notify.publish_after(con, {"type": "refresh", "scope": "products"})
    stocksync.wake()
    return {"counted": len(counts), "changed": changed, "same": same}


def preview(counts: dict[int, int]) -> dict:
    """What an upload would change, before anything is saved."""
    with db.read() as con:
        rows = db.rows(con.execute(f"""SELECT id, name, stock, woo_id, woo_stock, woo_state FROM products
                                       WHERE id IN ({','.join('?' * len(counts)) or 'NULL'})""", list(counts)))
    changes = [{"id": r["id"], "name": r["name"], "panel": r["stock"], "counted": counts[r["id"]],
                "site": r["woo_stock"] if r["woo_id"] and r["woo_state"] != "nocount" else None}
               for r in rows if counts[r["id"]] != r["stock"]]
    site = sum(1 for r in rows if r["woo_id"] and r["woo_state"] not in ("whole", "gone"))
    return {"counted": len(counts), "changed": len(changes), "site": site,
            "changes": sorted(changes, key=lambda c: -abs(c["counted"] - c["panel"]))}
